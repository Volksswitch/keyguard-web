"""House style for the Keyguard Designer's reader-facing documents.

Matches Conversant AAC's product documents (its generate-product-overview.py and
apply-doc-style.py): Arial 12 body, Word-blue headings that are not bold, a navy title
block, a grey byline, and a centered grey footer with page numbers. Paragraph spacing
is the Conversant house values: prose 8 pt after, list items 4 pt after, table cells 0.

Text passed to the helpers may carry **bold** and *italic*. Straight quotes and
apostrophes are turned typographic, because the documents are curly throughout.

    from kg_docstyle import KgDoc
    d = KgDoc(footer='Keyguard Designer Product Overview')
    d.title_block('Keyguard Designer', 'Custom keyguards in your browser', 'October 26, 2026')
    d.h1('Why keyguards matter'); d.p('A keyguard is ...'); d.save(path)
"""
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = RGBColor(0x1F, 0x38, 0x64)    # title block
BLUE = RGBColor(0x2E, 0x74, 0xB5)    # Word-default heading blue
BLUE3 = RGBColor(0x4F, 0x81, 0xBD)
GREY = RGBColor(0x55, 0x55, 0x55)    # byline
FOOT = RGBColor(0x80, 0x80, 0x80)    # header and footer

HERE = os.path.dirname(os.path.abspath(__file__))
PROJECT = os.path.abspath(os.path.join(HERE, '..', '..'))
DOCUMENTS = os.path.join(PROJECT, 'Documents')


def smart(s):
    """Typographic quotes and apostrophes."""
    s = re.sub(r"(\w)'(\w)", '\\1\u2019\\2', s)          # don't, it's
    s = re.sub(r"(\s|^|\(|\u2014)'", '\\1\u2018', s)     # opening single
    s = s.replace("'", '\u2019')
    s = re.sub(r'(\s|^|\(|\u2014)"', '\\1\u201c', s)     # opening double
    return s.replace('"', '\u201d')


def _force_arial(style):
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    for a in list(rf.attrib):
        del rf.attrib[a]
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(a), 'Arial')


def _field(run, instr):
    for tag, attr in (('w:fldChar', 'begin'), ('w:instrText', None),
                      ('w:fldChar', 'separate'), ('w:t', None), ('w:fldChar', 'end')):
        el = OxmlElement(tag)
        if tag == 'w:fldChar':
            el.set(qn('w:fldCharType'), attr)
        elif tag == 'w:instrText':
            el.set(qn('xml:space'), 'preserve')
            el.text = instr
        else:
            el.text = '1'
        run._r.append(el)


class KgDoc:
    def __init__(self, footer, header=''):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
            setattr(sec, side, Inches(1))
        self.sec = sec
        self._styles()
        self._header_footer(header, footer)

    # ---------- styles ----------
    def _styles(self):
        st = self.doc.styles
        normal = st['Normal']
        _force_arial(normal)
        normal.font.size = Pt(12)
        pf = normal.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(8), 1.08
        for name, size, color, before, after in (('Heading 1', 16, BLUE, 18, 6),
                                                 ('Heading 2', 13, BLUE, 14, 4),
                                                 ('Heading 3', 12, BLUE3, 10, 3)):
            h = st[name]
            _force_arial(h)
            h.font.size, h.font.bold, h.font.italic = Pt(size), False, False
            h.font.color.rgb = color
            h.paragraph_format.space_before = Pt(before)
            h.paragraph_format.space_after = Pt(after)
            h.paragraph_format.keep_with_next = True
        for name in ('List Bullet', 'List Number', 'List Bullet 2'):
            s = st[name]
            _force_arial(s)
            s.font.size = Pt(12)
            s.paragraph_format.space_before = Pt(0)
            s.paragraph_format.space_after = Pt(4)
            s.paragraph_format.line_spacing = 1.08

    def _header_footer(self, header, footer):
        if header:
            hp = self.sec.header.paragraphs[0]
            hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = hp.add_run(header)
            r.italic, r.font.size, r.font.color.rgb = True, Pt(9), FOOT
        fp = self.sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

        def run(text=''):
            r = fp.add_run(text)
            r.font.size, r.font.color.rgb = Pt(9), FOOT
            return r
        run(footer + '  |  Page ')
        _field(run(), 'PAGE')
        run(' of ')
        _field(run(), 'NUMPAGES')

    # ---------- text ----------
    def runs(self, p, text, size=None, color=None, bold_all=False, italic_all=False):
        for part in re.split(r'(\*\*.+?\*\*|(?<![\w*])\*[^*\s][^*]*?\*(?![\w*]))', smart(text)):
            if not part:
                continue
            bold = part.startswith('**') and part.endswith('**') and len(part) > 4
            italic = not bold and part.startswith('*') and part.endswith('*') and len(part) > 2
            r = p.add_run(part[2:-2] if bold else part[1:-1] if italic else part)
            r.bold = True if (bold or bold_all) else None
            r.italic = True if (italic or italic_all) else None
            if size:
                r.font.size = Pt(size)
            if color is not None:
                r.font.color.rgb = color
        return p

    def p(self, text, style=None):
        return self.runs(self.doc.add_paragraph(style=style), text)

    def b(self, text):
        return self.p(text, 'List Bullet')

    def b2(self, text):
        return self.p(text, 'List Bullet 2')

    def n(self, text):
        return self.p(text, 'List Number')

    def _heading(self, text, level):
        h = self.doc.add_heading(smart(text), level=level)
        h.paragraph_format.keep_with_next = True    # set on the paragraph too: Conversant's rule S3
        return h

    def h1(self, text):
        return self._heading(text, 1)

    def h2(self, text):
        return self._heading(text, 2)

    def h3(self, text):
        return self._heading(text, 3)

    def centered(self, text, size, color, bold=False, after=0, italic=False):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(after)
        self.runs(p, text, size=size, color=color, bold_all=bold, italic_all=italic)
        return p

    def title_block(self, title, subtitle, byline):
        self.centered(title, 32, NAVY, bold=True, after=2)
        self.centered(subtitle, 18, NAVY, after=8)
        self.centered(byline, 11, GREY, after=14)

    def quote(self, text, who):
        """An indented quotation with its attribution on the line below, as one paragraph."""
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.right_indent = Inches(0.5)
        self.runs(p, text, italic_all=True)
        p.runs[-1].add_break()
        self.runs(p, '\u2014 ' + who, size=10, color=GREY)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---------- tables ----------
    @staticmethod
    def _fill(cell, hexfill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), hexfill)
        tcPr.append(shd)

    @staticmethod
    def _borders(table, color='auto', size='8'):
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            el = OxmlElement('w:' + edge)
            el.set(qn('w:val'), 'single')
            el.set(qn('w:sz'), size)
            el.set(qn('w:space'), '0')
            el.set(qn('w:color'), color)
            borders.append(el)
        table._tbl.tblPr.append(borders)

    @staticmethod
    def _margins(table, top=60, bottom=60, left=100, right=100):
        mar = OxmlElement('w:tblCellMar')
        for k, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
            el = OxmlElement('w:' + k)
            el.set(qn('w:w'), str(v))
            el.set(qn('w:type'), 'dxa')
            mar.append(el)
        table._tbl.tblPr.append(mar)

    def _cell(self, cell, text, bold=False, size=11):
        cell.paragraphs[0].text = ''
        for i, line in enumerate(text.split('\n')):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            pf = p.paragraph_format
            pf.space_after, pf.space_before, pf.line_spacing = Pt(0), Pt(0), 1.0
            self.runs(p, line, size=size, bold_all=bold)

    def _widths(self, t, widths):
        for row in t.rows:
            for i, w in enumerate(widths):
                if i < len(row.cells):
                    row.cells[i].width = Inches(w)
        for i, gc in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))):
            gc.set(qn('w:w'), str(int(widths[i] * 1440)))

    def table(self, headers, rows, widths, first_col=True, size=11):
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        self._borders(t)
        self._margins(t)
        trPr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement('w:tblHeader')
        th.set(qn('w:val'), 'true')
        trPr.append(th)
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]
            self._cell(c, h, bold=True, size=size)
            self._fill(c, 'D9E2F3')
            c.paragraphs[0].paragraph_format.keep_with_next = True
        for row in rows:
            r = t.add_row()
            cant = OxmlElement('w:cantSplit')
            cant.set(qn('w:val'), 'true')
            r._tr.get_or_add_trPr().append(cant)
            for i, val in enumerate(row):
                c = r.cells[i]
                self._cell(c, val, bold=(first_col and i == 0), size=size)
                if first_col and i == 0:
                    self._fill(c, 'F0F4F8')
        self._widths(t, widths)
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)
        return t

    def callout(self, text, label='Tip', fill='EAF1FB', border='9DB9E0'):
        """A shaded one-cell box for a tip or a caution."""
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        self._borders(t, border, '6')
        self._margins(t, 100, 100, 160, 160)
        c = t.rows[0].cells[0]
        self._fill(c, fill)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        self.runs(p, f'**{label}.** {text}', size=11)
        self._widths(t, [6.5])
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)
        return t

    # ---------- pictures ----------
    def figure(self, path, caption, width=6.0):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(path, width=Inches(width))
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(10)
        self.runs(c, caption, size=10, color=GREY, italic_all=True)

    # ---------- contents ----------
    def contents(self, levels='1-2'):
        """A Word table of contents; Word fills it when the document is opened and updated."""
        self.h1('Contents')
        p = self.doc.add_paragraph()
        r = p.add_run()
        _field(r, f'TOC \\o "{levels}" \\h \\z \\u')

    def save(self, path):
        os.makedirs(os.path.dirname(path), exist_ok=True)
        self.doc.save(path)
        return path
