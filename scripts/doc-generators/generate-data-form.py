"""Builds Documents/Keyguard Data Form.docx: the form a clinician, a family or the person
who uses the tablet fills in, so that whoever designs the keyguard has what it needs.

It replaces the older form on volksswitch.org (Keyguard/data form/keyguard-data-form.docx),
which listed Customizer settings by their own names, half of them the designer's choices
rather than anything a customer would know (Ken, 30 Oct 2026: "an updated keyguard data
form that aligns with the current Customizer parameters ... only those items that would
likely be gathered as input from a customer (SLP, OT, AT specialist, user, or caregiver)").
So each question is asked in plain words, and the setting it fills is named under it in
small grey type, for the designer. Settings that are the designer's to choose (edge slopes,
chamfers, tab sizes, frames, splitting) are left out, as are the printer's settings.

Rows shaded yellow are measurements of the app taken with a ruler, which a screenshot makes
unnecessary: the same convention as the older form, which also shaded the rows and columns
(they're easy to count, and are asked for plainly here).

Video 2.4 is made from this form's pictures (scripts/video/scripts/2.4-data-form.mjs).

⚠ A generator rewrites the whole document. Once Ken has edited the .docx by hand, the
.docx is the newer copy.

    python scripts/doc-generators/generate-data-form.py [output.docx]
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from kg_docstyle import DOCUMENTS, KgDoc, NAVY, GREY  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt  # noqa: E402

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(DOCUMENTS, 'Keyguard Data Form.docx')
UPDATED = 'October 30, 2026'
BOX = '☐'
SIZE = 10.5
HINT = 8.5
WIDTHS = [3.45, 3.65]
HEAD_FILL, SHADE = 'D9E2F3', 'FFF2CC'

d = KgDoc(footer='Volksswitch.org  |  Keyguard Data Form', body_pt=SIZE, margin=0.7)


def boxes(*options):
    return '    '.join(f'{BOX} {o}' for o in options)


def cant_split(row):
    el = OxmlElement('w:cantSplit')
    el.set(qn('w:val'), 'true')
    row._tr.get_or_add_trPr().append(el)


def fill_cell(cell, text, hint=None, bold=False, italic=False, size=SIZE, color=None):
    p = cell.paragraphs[0]
    p.text = ''
    pf = p.paragraph_format
    pf.space_after, pf.space_before, pf.line_spacing = Pt(0), Pt(0), 1.0
    d.runs(p, text, size=size, bold_all=bold, italic_all=italic, color=color)
    if hint:
        h = cell.add_paragraph()
        h.paragraph_format.space_after, h.paragraph_format.space_before = Pt(0), Pt(1)
        d.runs(h, hint, size=HINT, color=GREY, italic_all=True)


def section(title, customizer, rows):
    """A table for one part of the form. Each row is (question, setting, answer[, shaded]);
    an answer of None makes the question a note across the row."""
    t = d.doc.add_table(rows=1, cols=2)
    t.style = d.doc.styles['Table Grid']
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    d._borders(t, '8EA9C9', '6')
    d._margins(t, 50, 50, 90, 90)
    head = t.rows[0]
    cant_split(head)
    fill_cell(head.cells[0], title, bold=True, size=SIZE + 1, color=NAVY)
    fill_cell(head.cells[1], f'Customizer: {customizer}' if customizer else '', italic=True,
              size=HINT + 0.5, color=GREY)
    for c in head.cells:
        d._fill(c, HEAD_FILL)
        c.paragraphs[0].paragraph_format.keep_with_next = True
    for question, setting, answer, *shaded in rows:
        r = t.add_row()
        cant_split(r)
        if answer is None:              # a note across the whole row
            fill_cell(r.cells[0].merge(r.cells[1]), question, italic=True, size=SIZE - 1)
            continue
        fill_cell(r.cells[0], question, hint=setting)
        fill_cell(r.cells[1], answer)
        if shaded and shaded[0]:
            for c in r.cells:
                d._fill(c, SHADE)
    d._widths(t, WIDTHS)
    sp = d.doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)
    return t


d.centered('Keyguard Data Form', 22, NAVY, bold=True, after=1)
d.centered(f'For a custom keyguard from the Keyguard Designer  |  Volksswitch.org  |  Last updated {UPDATED}',
           9, GREY, after=6)
d.p("Fill in what you know, and leave the rest blank. Whoever designs the keyguard enters these in the "
    "Keyguard Designer. Measure in millimeters. The most useful thing you can send is a **screenshot** of "
    "the app, taken on the tablet: with one, the rows shaded yellow can be skipped.")

section('About this keyguard', 'the preset', [
    ("A name for this design, such as the student's initials, the app and the tablet", "the preset's name", ''),
    ('Your name and role (SLP, OT, AT specialist, user, caregiver)', None, ''),
    ('Email or phone', None, ''),
    ('Date', None, ''),
])

section('The tablet', 'Tablet', [
    ('Make and model, with its generation (on an iPad: Settings, General, About)', 'type of tablet', ''),
    ('How the student holds it', 'orientation', boxes('Upright (portrait)', 'Sideways (landscape)')),
    ('Where the camera is, as it\'s held', 'rotate tablet 180 degrees',
     boxes('Top', 'Bottom', 'Left', 'Right')),
    ('Leave the home button uncovered?', 'expose home button', boxes('Yes', 'No', 'It has none')),
    ('Leave the camera uncovered?', 'expose camera', boxes('Yes', 'No')),
])

section('The case', 'Tablet Case', [
    ('Is the tablet in a case? Its make and model', 'have a case', f'{BOX} Yes    {BOX} No\nmake and model:'),
    ('The opening in the case, around the screen (mm)',
     'height of opening in case, width of opening in case, case opening corner radius',
     'height ______   width ______   corners ______'),
    ('How far the screen sits below the front of the case (mm)', 'case to screen depth', ''),
    ('Is the screen in the middle of the opening? If not, the gap from the opening\'s edge to the screen (mm)',
     'unequal left side of case opening, unequal bottom side of case opening (Special Actions and Settings)',
     f'{BOX} Yes    {BOX} No\nleft ______   bottom ______'),
    ('The outside of the case, for straps or clips (mm)', 'case height, case width, case thickness, case corner radius',
     'height ______   width ______\nthickness ______   corners ______'),
])

d.page_break()
section('The AAC app', 'App Layout, Bar Info, Grid Info, Grid Special Settings', [
    ('The app, and the vocabulary or page set', None, ''),
    ('A screenshot of each page that needs a keyguard, taken on the tablet held the way the student '
     'holds it, and sent full size', 'include screenshot; the designer measures App Layout in px from it',
     f'{BOX} Attached'),
    ('Rows and columns of buttons', 'number of rows, number of columns', 'rows ______   columns ______'),
    ('Bars the student needs to reach, so they get an opening. The message bar shows the message as it\'s '
     'built; a command bar is a row of buttons such as Speak and Clear.',
     'expose status bar, expose upper message bar, expose upper command bar, expose lower message bar, '
     'expose lower command bar',
     f'{BOX} Status bar\n{BOX} Upper message bar    {BOX} Upper command bar\n'
     f'{BOX} Lower message bar    {BOX} Lower command bar'),
    ('Buttons to cover, so they can\'t be pressed', 'cover these cells', ''),
    ('Buttons that are two cells big, across or down', 'merge cells horizontally starting at, '
     'merge cells vertically starting at', ''),
    ('Buttons to outline with a ridge, to find by touch', 'add a ridge around these cells', ''),
    ('Mark the last three on the screenshot, or list them by number, counting from the bottom-left '
     'button along the bottom row, then up the rows.', None, None),
    ('Without a screenshot: how tall each bar is (mm), with a ruler on the screen, 0 for a bar the app '
     'doesn\'t have', 'App Layout in mm', 'status ______   upper message ______\nupper command ______\n'
     'lower message ______   lower command ______', True),
    ('Without a screenshot: the size of one button (mm)', 'cell height in mm, cell width in mm',
     'height ______   width ______', True),
])

section('How the student uses it', 'Keyguard Basics, Grid Special Settings', [
    ('How the student touches the screen', 'keyguard thickness, cell edge slope',
     f'{BOX} One finger    {BOX} Several fingers\n{BOX} Knuckle or fist    {BOX} Stylus or pointer\n'
     f'{BOX} Other: ____________'),
    ('Anything else about how the student uses the tablet: rests a hand on it, slides a finger across it, '
     'presses hard', 'keyguard thickness, cell edge slope', '\n\n'),
])

d.page_break()
section('Making and attaching it', 'Keyguard Basics, Mounting Method', [
    ('Made by', 'type of keyguard',
     f'{BOX} 3D printer (plastic)    {BOX} Laser cutter (acrylic)\n{BOX} No preference'),
    ('How it should attach', 'mounting method',
     f'{BOX} Not sure: suggest one\n{BOX} Fits snugly in the case opening\n'
     f'{BOX} Velcro    {BOX} Suction cups    {BOX} Clip-on straps\n'
     f'{BOX} Screw-on straps    {BOX} Slide-in tabs    {BOX} Raised tabs\n'
     f'{BOX} Posts    {BOX} Shelf'),
    ('A favorite color', "not a setting: the plastic or acrylic it's made from", ''),
])

section('Words and braille', 'Engraved/Embossed Text, Cell Inserts', [
    ('Words to engrave, such as the student\'s name, and where', 'text, keyguard location',
     f'\n{BOX} On the top    {BOX} On the back'),
    ('Braille labels: which buttons, and the words', 'Braille text', '\n'),
])

section('Anything else', None, [
    ('Anything the designer should know', None, '\n' * 9),
])
d.p('Send the form, with the screenshots, to whoever is designing the keyguard.')

print(d.save(OUT))
