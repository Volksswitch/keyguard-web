#!/usr/bin/env python3
"""build-security-docx.py - render SECURITY.md to SECURITY.docx.

Trigger phrase: "build the security document".

WHY THIS EXISTS
---------------
The security document on volksswitch.org is the only copy anybody actually reads:
IT departments are handed it from the FAQ page. Nothing in the app links to the
markdown, and a school's IT team does not go to a git repository. So the published
file - not the markdown - is the deliverable.

Ken converts the .docx to PDF himself and uploads it, so this script stops at the
.docx. That is deliberate: he keeps editorial control of the final document, and an
earlier version of this tool that produced the PDF directly took that away.

THE DOCX IS AN OUTPUT. Never hand-edit it - edit SECURITY.md and rebuild, or the
next rebuild silently discards the change.

WHENEVER app.html changes what the app does on the network, in the browser's
storage, or with the clinician's files, update SECURITY.md and rebuild in the same
change. Between May and August 2026 the published copy drifted three months behind
the app it described - it still promised IT teams that the app contacted no third
party, months after the designer update check had shipped. A security document that
fails the test it invites the reader to run costs more trust than having none.

WordPress renames a re-uploaded file (SECURITY-1.pdf, or a new month folder), which
would break any link pointing at it and leave the stale copy live at the old
address. Ken replaces the file IN PLACE via the Media Library so the existing link
keeps working - which is also why the app deliberately contains no link to it.

    python scripts/build-security-docx.py                 -> ./SECURITY.docx
    python scripts/build-security-docx.py in.md           -> ./in.docx
    python scripts/build-security-docx.py in.md out.docx

Renders a deliberately plain, printable document: this is read by people deciding
whether to trust the app, so it should look like a document, not like marketing.
Requires python-docx (already present on both machines).
"""

import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Inches, RGBColor
except ImportError:
    sys.exit("build-security-docx: python-docx is not installed.\n"
             "  Install it with:  python -m pip install python-docx")

ROOT = Path(__file__).resolve().parent.parent

# Defaults render SECURITY.md, which is what the trigger phrase does. Both paths
# are overridable so the same renderer can produce any of the project's markdown
# as Word - Ken reads .docx, and several documents here are written for him.
#     build-security-docx.py                       -> SECURITY.md  -> SECURITY.docx
#     build-security-docx.py <in.md>               -> <in.md>      -> <in>.docx
#     build-security-docx.py <in.md> <out.docx>    -> explicit both
SRC = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else ROOT / "SECURITY.md"
if len(sys.argv) > 2:
    OUT = Path(sys.argv[2]).resolve()
elif len(sys.argv) > 1:
    OUT = SRC.with_suffix(".docx")
else:
    OUT = ROOT / "SECURITY.docx"

if not SRC.exists():
    sys.exit(f"build-security-docx: {SRC} not found.")

# Inline spans, in precedence order. Code first so its contents stay literal.
INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|(?<!\*)\*[^*\s][^*]*\*(?!\*)|_[^_\s][^_]*_)")


def add_runs(par, text):
    """Write markdown inline formatting into a paragraph."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("`") and piece.endswith("`"):
            run = par.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0xA0, 0x30, 0x30)
        elif piece.startswith("**") and piece.endswith("**"):
            par.add_run(piece[2:-2]).bold = True
        elif piece.startswith("*") and piece.endswith("*"):
            par.add_run(piece[1:-1]).italic = True
        elif piece.startswith("_") and piece.endswith("_"):
            par.add_run(piece[1:-1]).italic = True
        else:
            par.add_run(piece)


def horizontal_rule(doc):
    """A thin section divider, standing in for markdown's ---."""
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(2)
    par.paragraph_format.space_after = Pt(8)
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "BBBBBB")
    borders.append(bottom)
    par._p.get_or_add_pPr().append(borders)


def shade(cell_or_par, hex_fill):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_fill)
    cell_or_par.append(el)


def page_number_footer(doc, label):
    """Page N in the footer - this document gets printed and passed around."""
    par = doc.sections[0].footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run(label + "    Page ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    # PAGE is a Word field; it has to be built out of raw XML.
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    par._p.append(fld)


def build():
    lines = SRC.read_text(encoding="utf-8").split("\n")
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(7)
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(1.0)
        section.top_margin = section.bottom_margin = Inches(0.9)

    counts = {"heading": 0, "paragraph": 0, "bullet": 0, "numbered": 0,
              "table": 0, "code": 0}
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # --- horizontal rule ---
        if stripped == "---":
            horizontal_rule(doc)
            i += 1
            continue

        # --- heading ---
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            if level == 1:
                par = doc.add_paragraph(style="Title")
            else:
                par = doc.add_paragraph(style=f"Heading {min(level - 1, 4)}")
            add_runs(par, m.group(2))
            counts["heading"] += 1
            i += 1
            continue

        # --- fenced code block ---
        if stripped.startswith("```"):
            i += 1
            body = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                body.append(lines[i])
                i += 1
            i += 1
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Inches(0.25)
            par.paragraph_format.space_before = Pt(4)
            run = par.add_run("\n".join(body))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            shade(par._p.get_or_add_pPr(), "F4F4F4")
            counts["code"] += 1
            continue

        # --- table ---
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            # Row 2 of a markdown table is the |---|---| separator.
            if len(rows) >= 2 and all(set(c) <= set("-: ") for c in rows[1]):
                header, body = rows[0], rows[2:]
            else:
                header, body = rows[0], rows[1:]
            table = doc.add_table(rows=0, cols=len(header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cells = table.add_row().cells
            for cell, text in zip(cells, header):
                cell.paragraphs[0].text = ""
                add_runs(cell.paragraphs[0], text)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                shade(cell._tc.get_or_add_tcPr(), "EFEFEF")
            for row in body:
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    cell.paragraphs[0].text = ""
                    add_runs(cell.paragraphs[0], text)
            doc.add_paragraph()
            counts["table"] += 1
            continue

        # --- numbered list (ordering is content, so it gets a real Word list) ---
        m = re.match(r"^(\s*)\d+\.\s+(.*)$", line)
        if m:
            indent = len(m.group(1))
            text = [m.group(2)]
            i += 1
            while i < len(lines):
                nxt = lines[i]
                if (not nxt.strip() or re.match(r"^\s*(\d+\.|[-*])\s+", nxt)
                        or nxt.startswith("#")):
                    break
                text.append(nxt.strip())
                i += 1
            par = doc.add_paragraph(style="List Number" if indent < 2 else "List Number 2")
            add_runs(par, " ".join(text))
            counts["numbered"] += 1
            continue

        # --- bullet (with continuation lines, and one nesting level) ---
        m = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if m:
            indent = len(m.group(1))
            text = [m.group(2)]
            i += 1
            while i < len(lines):
                nxt = lines[i]
                if (not nxt.strip() or re.match(r"^\s*(\d+\.|[-*])\s+", nxt)
                        or nxt.startswith("#")):
                    break
                text.append(nxt.strip())
                i += 1
            style_name = "List Bullet" if indent < 2 else "List Bullet 2"
            par = doc.add_paragraph(style=style_name)
            add_runs(par, " ".join(text))
            counts["bullet"] += 1
            continue

        # --- blockquote ---
        if stripped.startswith(">"):
            text = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                text.append(lines[i].strip().lstrip(">").strip())
                i += 1
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Inches(0.3)
            add_runs(par, " ".join(t for t in text if t))
            for run in par.runs:
                run.italic = True
            counts["paragraph"] += 1
            continue

        # --- paragraph (joining wrapped lines) ---
        text = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            if (not nxt.strip() or nxt.startswith("#") or nxt.strip() == "---"
                    or nxt.strip().startswith(("|", "```", ">"))
                    or re.match(r"^\s*(\d+\.|[-*])\s+", nxt)):
                break
            text.append(nxt.strip())
            i += 1
        par = doc.add_paragraph()
        add_runs(par, " ".join(text))
        counts["paragraph"] += 1

    page_number_footer(doc, "Keyguard Designer - Security & Privacy")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return counts


counts = build()
print(f"Wrote {OUT}")
print("  " + ", ".join(f"{v} {k}{'s' if v != 1 else ''}" for k, v in counts.items()))
if SRC.name == "SECURITY.md":
    print("\nConvert to PDF yourself and replace the file IN PLACE in the WordPress Media")
    print("Library, so the existing link keeps working.")
